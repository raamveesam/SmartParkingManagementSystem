import os, subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

REPO = r"C:\Users\RamLaxma\.gemini\antigravity\scratch\SmartParkingManagementSystem"
OUT  = os.path.join(REPO, "Smart_Parking_Management_System_Code.docx")

def code_block(doc, text):
    tbl  = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    # light-gray background
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>'))
    # blue left border
    tcPr.append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top    w:val="single" w:sz="4"  w:color="CBD5E1"/>'
        f'<w:left   w:val="single" w:sz="18" w:color="1D4ED8"/>'
        f'<w:bottom w:val="single" w:sz="4"  w:color="CBD5E1"/>'
        f'<w:right  w:val="single" w:sz="4"  w:color="CBD5E1"/>'
        f'</w:tcBorders>'))
    # inner margin
    tcPr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="100" w:type="dxa"/><w:bottom w:w="100" w:type="dxa"/>'
        f'<w:left w:w="160" w:type="dxa"/><w:right w:w="160" w:type="dxa"/>'
        f'</w:tcMar>'))
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(8)

def main():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1.0)
        sec.left_margin = sec.right_margin  = Inches(1.0)

    # ---- Title ----
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.space_before = Pt(20)
    r = tp.add_run("BAYTRACK — SMART PARKING MANAGEMENT SYSTEM")
    r.font.name = "Arial"; r.font.size = Pt(20)
    r.font.bold = True; r.font.color.rgb = RGBColor(30, 58, 138)

    sp = doc.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp.paragraph_format.space_after = Pt(16)
    rs = sp.add_run(
        "Data Structures & AI Project in C (C99)\n"
        "GitHub: github.com/raamveesam/SmartParkingManagementSystem")
    rs.font.name = "Arial"; rs.font.size = Pt(11)
    rs.font.italic = True; rs.font.color.rgb = RGBColor(100, 116, 139)

    # ---- Abstract ----
    h = doc.add_heading("1. Project Abstract & DSA Overview", level=1)
    h.paragraph_format.space_before = Pt(12)
    body = doc.add_paragraph(
        "BayTrack is a multi-level parking management system written entirely in C99. "
        "It demonstrates the following data structures and algorithms:\n\n"
        "  • Static Array  (bay_grid[])  — physical bay storage & slot allocation\n"
        "  • Singly Linked List (ExitLog) — prepend-based billing audit trail\n"
        "  • FIFO Queue (HoldLine)        — overflow vehicle waiting line\n"
        "  • Linear Search               — locate a parked vehicle in bay_grid[]\n"
        "  • Insertion Sort              — rank completed sessions by fee collected\n"
        "  • Binary Search               — O(log n) pass-number lookup\n"
        "  • Gaussian AI Heuristic       — PeakScore forecasts hourly demand; "
        "pricing_mult() maps demand to a dynamic tariff (×0.80 off-peak → ×1.60 peak surge)\n\n"
        "Tariff: Two-Wheelers Rs.20/hr | Four-Wheelers Rs.50/hr (Indian Rupees)\n"
        "Demo uses real Indian vehicle registration formats (KA, TN, MH, AP …).")
    body.paragraph_format.line_spacing = 1.2
    body.paragraph_format.space_after  = Pt(14)

    # ---- Source code ----
    doc.add_heading("2. Complete Source Code  (main.c)", level=1).paragraph_format.space_before = Pt(14)
    with open(os.path.join(REPO, "main.c"), encoding="utf-8") as f:
        code_block(doc, f.read())

    # ---- Execution output ----
    doc.add_heading("3. Automated Demo Execution Output", level=1).paragraph_format.space_before = Pt(14)
    exe = os.path.join(REPO, "smart_parking.exe")
    out = ""
    if os.path.exists(exe):
        res = subprocess.run([exe, "--demo"], capture_output=True, text=True)
        out = res.stdout
    code_block(doc, out if out else "(executable not found)")

    doc.save(OUT)
    print(f"Saved: {OUT}")

if __name__ == "__main__":
    main()
